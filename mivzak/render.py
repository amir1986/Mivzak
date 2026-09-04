"""Render the brief into the IBI Word template and into the email body."""

from __future__ import annotations

import html
import re
from copy import deepcopy
from datetime import date
from pathlib import Path

from .config import CLOSING, OPENING, PAUSE, TEMPLATE_PATH, hebrew_full_date

_LATIN_RUN_RE = re.compile(r"[A-Za-z][A-Za-z0-9&.'\-]*(?:\s+[A-Za-z0-9&.'\-]+)*")


def split_runs(text: str) -> list:
    """Split text into (segment, is_latin) so Word keeps Latin tokens LTR."""
    segments = []
    position = 0
    for match in _LATIN_RUN_RE.finditer(text):
        if match.start() > position:
            segments.append((text[position:match.start()], False))
        segments.append((match.group(0), True))
        position = match.end()
    if position < len(text):
        segments.append((text[position:], False))
    return [segment for segment in segments if segment[0]]


_W14_NS = "http://schemas.microsoft.com/office/word/2010/wordml"


def _strip_paragraph_ids(paragraph_element) -> None:
    """Copied paragraphs must not share the template's unique ids."""
    for attribute in (f"{{{_W14_NS}}}paraId", f"{{{_W14_NS}}}textId"):
        if attribute in paragraph_element.attrib:
            del paragraph_element.attrib[attribute]


def _clear_runs(paragraph_element) -> None:
    from docx.oxml.ns import qn

    _strip_paragraph_ids(paragraph_element)
    for child in list(paragraph_element):
        if child.tag != qn("w:pPr"):
            paragraph_element.remove(child)


def _run_properties(paragraph_element, want_rtl: bool):
    """Copy run properties from the template, with or without <w:rtl/>."""
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    chosen = None
    for run in paragraph_element.findall(qn("w:r")):
        rpr = run.find(qn("w:rPr"))
        if rpr is None:
            continue
        has_rtl = rpr.find(qn("w:rtl")) is not None
        if has_rtl == want_rtl:
            chosen = deepcopy(rpr)
            break
        if chosen is None:
            chosen = deepcopy(rpr)
    if chosen is None:
        chosen = OxmlElement("w:rPr")
        fonts = OxmlElement("w:rFonts")
        for attribute in ("w:ascii", "w:hAnsi", "w:cs"):
            fonts.set(qn(attribute), "Arial")
        chosen.append(fonts)
    rtl = chosen.find(qn("w:rtl"))
    if want_rtl and rtl is None:
        chosen.append(OxmlElement("w:rtl"))
    if not want_rtl and rtl is not None:
        chosen.remove(rtl)
    hint = chosen.find(qn("w:rFonts"))
    if hint is not None and hint.get(qn("w:hint")) is not None:
        del hint.attrib[qn("w:hint")]
    return chosen


def _append_text_paragraph(cell_element, prototype, hebrew_rpr, latin_rpr, text: str) -> None:
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    paragraph = deepcopy(prototype)
    _clear_runs(paragraph)
    for segment, is_latin in split_runs(text):
        run = OxmlElement("w:r")
        run.append(deepcopy(latin_rpr if is_latin else hebrew_rpr))
        text_element = OxmlElement("w:t")
        text_element.text = segment
        text_element.set(qn("xml:space"), "preserve")
        run.append(text_element)
        paragraph.append(run)
    cell_element.append(paragraph)


def find_body_cell(document):
    if not document.tables:
        raise RuntimeError("The template has no table")
    table = document.tables[0]
    candidates = [cell for row in table.rows for cell in row.cells]
    for cell in candidates:
        if PAUSE in cell.text:
            return cell
    return max(candidates, key=lambda cell: len(cell.text))


def render_docx(paragraph_texts: list, brief_date: date, output_path: Path,
                template_path: Path = TEMPLATE_PATH) -> Path:
    from docx import Document

    if not Path(template_path).exists():
        raise RuntimeError(f"Template is missing: {template_path}")
    document = Document(str(template_path))
    body_cell = find_body_cell(document)
    cell_element = body_cell._tc

    text_prototype = None
    pause_prototype = None
    empty_prototype = None
    for paragraph in body_cell.paragraphs:
        content = paragraph.text.strip()
        if PAUSE in content:
            pause_prototype = paragraph._p
        elif content and text_prototype is None:
            text_prototype = paragraph._p
        elif not content and empty_prototype is None:
            empty_prototype = paragraph._p
    if text_prototype is None or pause_prototype is None:
        raise RuntimeError("The template body cell does not look like the IBI brief")
    if empty_prototype is None:
        empty_prototype = deepcopy(text_prototype)
        _clear_runs(empty_prototype)

    hebrew_rpr = _run_properties(text_prototype, want_rtl=True)
    latin_rpr = _run_properties(text_prototype, want_rtl=False)

    for paragraph in list(body_cell.paragraphs):
        cell_element.remove(paragraph._p)
    for text in paragraph_texts:
        lines = [line.strip() for line in str(text).split("\n") if line.strip()]
        for line in lines:
            _append_text_paragraph(cell_element, text_prototype, hebrew_rpr, latin_rpr, line)
        spacer = deepcopy(empty_prototype)
        _strip_paragraph_ids(spacer)
        cell_element.append(spacer)
    cell_element.append(pause_prototype)

    from datetime import datetime, timezone

    document.core_properties.title = "מבזק הידע של IBI על הבוקר " + brief_date.strftime("%d.%m.%Y")
    document.core_properties.subject = "מבזק שוק הון יומי לקריינות"
    document.core_properties.modified = datetime.now(timezone.utc).replace(tzinfo=None)
    output_path = Path(output_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    document.save(str(output_path))
    return output_path


def docx_body_text(path: Path) -> list:
    """Paragraph texts of the body cell (used by tests and the self-test)."""
    from docx import Document

    document = Document(str(path))
    return [paragraph.text for paragraph in find_body_cell(document).paragraphs if paragraph.text.strip()]


def collect_sources(paragraphs: list) -> list:
    seen = set()
    result = []
    for paragraph in paragraphs:
        for source in paragraph.sources:
            key = source.get("url")
            if not key or key in seen:
                continue
            seen.add(key)
            result.append(source)
    return result


def email_subject(brief_date: date) -> str:
    return "מבזק בוקר " + brief_date.strftime("%d.%m.%Y")


def docx_filename(brief_date: date) -> str:
    return "מבזק בוקר " + brief_date.strftime("%d.%m.%Y") + ".docx"


def paragraph_lines(paragraph) -> list:
    text = paragraph.text if hasattr(paragraph, "text") else str(paragraph)
    return [line.strip() for line in text.split("\n") if line.strip()]


def render_email(trading_date: date, brief_date: date, paragraphs: list, notes: list) -> tuple:
    """Return (plain text, HTML) mirroring the template table."""
    sources = collect_sources(paragraphs)
    plain_lines = [OPENING, ""]
    for paragraph in paragraphs:
        plain_lines.extend(paragraph_lines(paragraph) + [""])
    plain_lines.extend([PAUSE, "", CLOSING, ""])
    plain_lines.append("כל הנתונים במבזק מתייחסים ליום המסחר " + trading_date.strftime("%d.%m.%Y") + " בלבד.")
    if sources:
        plain_lines.extend(["", "מקורות שנבדקו:"])
        plain_lines.extend(f"{source['label']}: {source['url']}" for source in sources)
    if notes:
        plain_lines.extend(["", "הערות מערכת:"])
        plain_lines.extend(f"- {note}" for note in notes)

    rich_paragraphs = "".join(
        '<p style="margin:0 0 16px 0;">'
        + "<br>".join(html.escape(line) for line in paragraph_lines(paragraph))
        + "</p>"
        for paragraph in paragraphs
    )
    rich_sources = "".join(
        '<li style="margin:0 0 6px 0;"><a href="' + html.escape(source["url"], quote=True)
        + '" style="color:#1155cc;">' + html.escape(f"{source['label']} — {source.get('title') or source['url']}")
        + "</a></li>"
        for source in sources
    )
    rich_notes = "".join("<li>" + html.escape(note) + "</li>" for note in notes)
    notes_block = (
        '<p style="margin:18px 0 6px 0;font-size:14px;"><strong>הערות מערכת</strong></p>'
        f'<ul style="margin:0;padding-right:22px;font-size:13px;color:#555;">{rich_notes}</ul>'
        if notes
        else ""
    )
    rich = f"""<!doctype html>
<html lang="he" dir="rtl">
  <head><meta charset="UTF-8"></head>
  <body dir="rtl" style="direction:rtl;text-align:right;background:#ffffff;color:#202124;font-family:Arial,Helvetica,sans-serif;font-size:17px;line-height:1.65;margin:0;">
    <div style="max-width:820px;margin:0 auto;padding:24px;">
      <table role="presentation" style="width:100%;border-collapse:collapse;border:1px solid #444;">
        <tr>
          <td style="padding:10px 14px;font-weight:bold;border:1px solid #444;">{html.escape(OPENING)}</td>
          <td style="width:74px;padding:10px 8px;text-align:center;font-weight:bold;border:1px solid #444;">פתיחה</td>
        </tr>
        <tr>
          <td style="padding:18px 14px;border:1px solid #444;vertical-align:top;">
            {rich_paragraphs}
            <p style="margin:4px 0 0 0;"><strong><span style="background:#00ff00;">{html.escape(PAUSE)}</span></strong></p>
          </td>
          <td style="border:1px solid #444;">&nbsp;</td>
        </tr>
        <tr>
          <td style="padding:10px 14px;font-weight:bold;border:1px solid #444;">{html.escape(CLOSING)}</td>
          <td style="padding:10px 8px;text-align:center;font-weight:bold;border:1px solid #444;">קריין</td>
        </tr>
      </table>
      <p style="margin:24px 0 8px 0;font-size:15px;"><strong>מקורות שנבדקו ואומתו ליום המסחר {trading_date.strftime('%d.%m.%Y')}</strong></p>
      <ul style="margin:0 0 18px 0;padding-right:22px;font-size:14px;">{rich_sources}</ul>
      {notes_block}
      <p style="margin:18px 0 0 0;font-size:13px;color:#666;">מצורף קובץ Word לפי תבנית המבזק, לבוקר {html.escape(hebrew_full_date(brief_date))}.</p>
    </div>
  </body>
</html>
"""
    return "\n".join(plain_lines), rich
